"""Fill the official Apart template (October 2026) from report.md and export PDF.

Written during the sprint. Usage:
  python report-collusion/build_docx.py            # writes submission-collusion/report.docx and report.pdf

The template's guidance paragraphs are removed and replaced with the report
body. Markdown handled: '## ' and '### ' headings, paragraphs, '> ' quotes,
pipe tables, '![...](path)' images, **bold** and `code` inline, '**Author:**' line.
"""

from __future__ import annotations

import copy
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "official-template-2026-10.docx"
SRC = ROOT / "report-collusion" / "report.md"
OUT_DIR = ROOT / "submission-collusion"
OUT_DOCX = OUT_DIR / "report.docx"
OUT_PDF = OUT_DIR / "report.pdf"


def add_inline(paragraph, text: str, size: float | None = None, italic: bool = False) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1.5)
        else:
            run = paragraph.add_run(part)
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)


def clear_body_after_tables(doc: Document) -> None:
    """Remove every body paragraph and table after the title table."""
    body = doc.element.body
    first_table = doc.tables[0]._tbl
    seen = False
    for child in list(body):
        if child is first_table:
            seen = True
            continue
        if not seen:
            continue
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_cell_paragraph(paragraph, text: str, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
        if bold is not None:
            paragraph.runs[0].bold = bold
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def set_table_borders(table) -> None:
    """Thin single borders on every cell (the template has no 'Table Grid' style)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for edge, width in (("top", "20"), ("bottom", "20"), ("left", "60"), ("right", "60")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), width)
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.append(margins)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            text = text.replace("\\", "")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, text, size=8.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def build() -> None:
    doc = Document(str(TEMPLATE))
    lines = SRC.read_text(encoding="utf-8").splitlines()

    title = next(l[2:] for l in lines if l.startswith("# "))
    author = next(l for l in lines if l.startswith("**Author:**"))
    author = re.sub(r"\*\*", "", author).replace("Author: ", "")
    abs_start = lines.index("## Abstract") + 1
    abs_end = next(i for i in range(abs_start, len(lines)) if lines[i].startswith("## "))
    abstract = " ".join(l for l in lines[abs_start:abs_end] if l.strip())

    title_cell = doc.tables[0].rows[0].cells[0]
    set_cell_paragraph(title_cell.paragraphs[1], title, bold=True)

    # Row 1 of the title table holds a nested author grid and a nested abstract box.
    abstract_cell = doc.tables[0].rows[1].cells[0]
    author_grid = abstract_cell.tables[0].cell(0, 0).tables[0]
    name, _, affiliation = author.partition(", ")
    affiliation = affiliation.split(". ")[0]
    author_cells = [c for row in author_grid.rows for c in row.cells]
    # Merge the six author slots into one full-width cell (solo submission).
    merged = author_cells[0].merge(author_cells[-1])
    for p in merged.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    author_cells = [merged]
    for k, cell in enumerate(author_cells):
        paras = cell.paragraphs
        for p in paras:
            set_cell_paragraph(p, "")
        if k == 0:
            run = paras[0].runs[0] if paras[0].runs else paras[0].add_run("")
            run.text = name
            run.add_break()
            paras[0].add_run(affiliation)
    abstract_box = abstract_cell.tables[1].cell(0, 0)
    set_cell_paragraph(abstract_box.paragraphs[0], abstract)
    for run in abstract_box.paragraphs[0].runs:
        run.italic = False
        run.underline = False
    for p in abstract_box.paragraphs[1:]:
        set_cell_paragraph(p, "")

    clear_body_after_tables(doc)

    i = abs_end
    table_buf: list[str] = []
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if line.startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            add_table(doc, parse_table(table_buf))
            table_buf = []
        if not line.strip():
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("!["):
            m = re.search(r"\]\(([^)]+)\)", line)
            path = (SRC.parent / m.group(1)).resolve()
            doc.add_picture(str(path), width=Inches(4.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, line[2:], size=9.5, italic=True)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            add_inline(p, line)
            p.paragraph_format.left_indent = Inches(0.25)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line)
    if table_buf:
        add_table(doc, parse_table(table_buf))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    print(f"wrote {OUT_DOCX}; abstract {len(abstract.split())} words")


def export_pdf() -> None:
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{OUT_DOCX}'); $d.ExportAsFixedFormat('{OUT_PDF}', 17); "
        "$pages = $d.ComputeStatistics(2); "
        "$refpage = -1; foreach ($p in $d.Paragraphs) { if ($p.Range.Text.Trim() -eq 'References') { $refpage = $p.Range.Information(3); break } }; "
        "$d.Close(0); $w.Quit(); Write-Output \"pages=$pages references_start_page=$refpage\""
    )
    out = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print(out.stdout.strip(), out.stderr.strip())


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        export_pdf()
